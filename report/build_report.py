"""Builds the project report (Word, LNCS template) from the results of the analysis.

Usage (from the repository root):  python report/build_report.py
Requires: python-docx, and report/template_lncs.docx (the course template).
"""
import copy
import json
import os
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = f"{ROOT}/report/template_lncs.docx"
OUT = f"{ROOT}/report/Report_Consumer_Choice_baseline.docx"
FIG = f"{ROOT}/results/figures"
TAB = f"{ROOT}/results/tables"
S = json.load(open(f"{TAB}/summary.json"))
TEXT_W = 122  # mm, LNCS printing area

doc = docx.Document(TEMPLATE)
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
KEEP = os.environ.get("KEEP", "1") == "1"
counters = {"table": 0, "fig": 0, "listing": 0}


# ---------------------------------------------------------------- helpers
def add_runs(par, text, size=None):
    """Minimal inline markup: *italic*, `code`, **bold**, ^sup^."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Courier New"
            r.font.size = Pt(size - 1 if size else 8.5)
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        elif tok.startswith("^"):
            r = par.add_run(tok[1:-1]); r.font.superscript = True
        else:
            r = par.add_run(tok)
        if size and not tok.startswith("`"):
            r.font.size = Pt(size)
    return par


_after_heading = [True]


def P(text, style=None):
    style = style or ("p1a" if _after_heading[0] else "Normal")
    _after_heading[0] = False
    return add_runs(doc.add_paragraph(style=style), text)


def H1(text):
    _after_heading[0] = True
    doc.add_paragraph(text, style="heading1")


def H2(text):
    _after_heading[0] = True
    doc.add_paragraph(text, style="heading2")


def H3(title, text):
    """Run-in heading (LNCS 3rd level): bold title followed by text."""
    p = doc.add_paragraph(style="heading3")
    r = p.add_run(title + " "); r.bold = True
    add_runs(p, text)
    for rr in p.runs[1:]:
        rr.bold = False
    _after_heading[0] = False
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-3.5)
        add_runs(p, "– " + it)


def caption(kind, text):
    counters[kind] += 1
    n = counters[kind]
    label = {"table": "Table", "fig": "Fig.", "listing": "Listing"}[kind]
    style = "table title" if kind == "table" else "figure legend"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = (kind == "table")
    r = p.add_run(f"{label} {n}. "); r.bold = True
    add_runs(p, text)
    return n


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, val in kw.items():
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:color"), "000000")
        b.append(e)


def table(cap, header, rows, widths, size=7.5, align=None):
    """LNCS-style table: caption above, rules at top, below header and at bottom."""
    caption("table", cap)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Mm(widths[j])
            p = c.paragraphs[0]
            p.style = doc.styles["Normal"]
            pf = p.paragraph_format
            pf.first_line_indent = Mm(0); pf.space_before = Pt(1); pf.space_after = Pt(1)
            pf.line_spacing = 1.0
            if align and align[j] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, str(val), size=size)
            if i == 0:
                for rr in p.runs:
                    rr.bold = True
            if i == 0:
                set_cell_border(c, top=8, bottom=4)
            if i == len(rows):
                set_cell_border(c, bottom=8)
    for j, w in enumerate(widths):
        for cell in t.columns[j].cells:
            cell.width = Mm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Mm(w).twips)))
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(Mm(sum(widths)).twips))); tw.set(qn("w:type"), "dxa")
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tw)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
             "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
             "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = sorted(list(tblPr), key=lambda e: order.index(e.tag.split("}")[1])
                  if e.tag.split("}")[1] in order else len(order))
    for k in kids:
        tblPr.remove(k)
        tblPr.append(k)
    doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(0)
    return t


def span(file, start, end):
    """Line numbers (1-based) from the first line containing `start`
    to the first following line containing `end`."""
    lines = open(f"{ROOT}/{file}").read().splitlines()
    a = next(i for i, l in enumerate(lines) if start in l)
    b = next(i for i in range(a, len(lines)) if end in lines[i])
    return (a + 1, b + 1)


def listing(file, segments, cap):
    """Code excerpt from the repository with its original line numbers."""
    lines = open(f"{ROOT}/{file}").read().splitlines()
    segments = [span(file, *seg) if isinstance(seg[0], str) else seg
                for seg in segments]
    p = doc.add_paragraph(style="programcode")
    p.paragraph_format.keep_together = KEEP
    first = True
    for k, (a, b) in enumerate(segments):
        if k:
            p.add_run().add_break(); p.add_run("   ...")
        for n in range(a, b + 1):
            if not first:
                p.add_run().add_break()
            first = False
            r = p.add_run(f"{n:>3}  {lines[n - 1]}")
            r.font.size = Pt(7)
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
            r._element.rPr.rFonts.set(qn("w:cs"), "Courier New")
    rng = ", ".join(f"{a}–{b}" for a, b in segments)
    caption("listing", f"{cap} (`{file}`, lines {rng}).")


def figure(path, cap, width=TEXT_W):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Mm(width))
    caption("fig", cap)


def f2(x):
    return f"{x:.2f}"


def f3(x):
    return f"{x:.3f}"


# ---------------------------------------------------------------- numbers
B = S["baseline"]
Bo, Bb, Bc = B["Orange market share"], B["Blue market share"], B["Consumer clustering"]
OUTC = S["baseline_outcomes"]
HZ = {r["tick"]: r for r in S["baseline_horizon"]}
CI100, CIALL = S["ci_halfwidth_at_100_runs"], S["ci_halfwidth_at_all_runs"]
SN = {r["norm_influence"]: r for r in S["sensitivity_norm_influence"]}
SQ = {r["quality_blue"]: r for r in S["sensitivity_quality_blue"]}
SE = {r["exploration"]: r for r in S["sensitivity_exploration"]}
RUNS = OUTC["runs"]

# ================================================================ front matter
doc.add_paragraph("Consumer Choice: Model Analysis and Baseline Simulation of an "
                  "Agent-Based Model of Habits and Social Influence", style="title")
doc.add_paragraph("[Author 1], [Author 2], [Author 3], [Author 4]", style="author")
doc.add_paragraph("[Institution], [Programme]", style="authorinfo")
doc.add_paragraph("{[student1], [student2], [student3], [student4]}@[domain]", style="email")

ab = doc.add_paragraph(style="abstract")
r = ab.add_run("Abstract."); r.bold = True
add_runs(ab, " This report analyses the Consumer Choice model, an agent-based model in which "
         "consumers choose between two products under the influence of habits, information "
         "from friends and social norms. The model was implemented in Python with Mesa 3, "
         "following the rules of the original NetLogo model. We describe the agents, their "
         f"state variables and actions, and define a baseline. We then run {RUNS} replications "
         "of the baseline and a sensitivity analysis of three parameters, using only the "
         "outputs of the model: the market shares of the two products and consumer "
         f"clustering. In the baseline, clustering reaches {f2(Bc['mean'])} at tick 1000, "
         "while the final market share of each product varies widely between runs "
         f"(standard deviation {f2(Bo['sd'])}). A small difference in quality determines "
         "which product dominates. These results define the reference for the extension "
         "of the model to a real market.")
kw = doc.add_paragraph(style="abstract")
r = kw.add_run("Keywords:"); r.bold = True
kw.add_run(" Agent-based modelling, Consumer behaviour, Social influence, Habits, Mesa.")

# ================================================================ 1 Introduction
H1("1   Introduction")
P("This report addresses the first two tasks of Project 1 of the course Multi-Agent Systems "
  "and Simulation of Organizations: (1) to analyse the chosen model and run it with a basic "
  "setup that becomes the baseline, and (2) to run the simulations for a large number of "
  "iterations to obtain meaningful statistics for the variables of interest. The remaining "
  "tasks (fitting the model to a real situation, learning agents and the analysis of the new "
  "model) will build on the baseline defined here.")
P("The chosen model is Consumer Choice [1], from the NetLogo Models Library [2]. It derives "
  "from a basic model by M. A. Janssen and from earlier work on habits and social processes in "
  "market dynamics [3]. It simulates "
  "a market with two competing products whose quality can change. Consumers are linked in a "
  "social network and choose a product according to three mechanisms: habit, informative "
  "social influence (learning about quality from friends) and normative social influence "
  "(using the same product as the majority of friends) [4]. The model shows how these "
  "mechanisms produce clusters of consumption, tipping points and social lock-in, a situation "
  "in which a product of lower quality keeps its market because it is the norm [5].")
P("The model was implemented in Python with the Mesa 3 framework [6]. The implementation "
  "follows the NetLogo code rule by rule. The description in Section 2 is organised along the "
  "lines of the ODD protocol [7] and places each rule next to the code that implements it. "
  "Section 3 presents the baseline experiments and their statistical analysis. Section 4 "
  "outlines how the model will be fitted to a real market in the next phase of the project. "
  "The code, data and figures are available in the project repository.")

# ================================================================ 2 Model
H1("2   Model Description")
H2("2.1   Purpose and Entities")
P("The purpose of the model is to explain how habits and social influence shape the market "
  "shares of two products, and why consumers may fail to adopt a product of higher quality. "
  "The model contains one type of agent, the consumer. The social network and the products are "
  "not agents (Table 1).")
table("Entities of the model.",
      ["Entity", "Represents", "Role in the model"],
      [["Consumer (100)", "A person who buys one of two products",
        "The only agent type. It chooses a product, learns about quality and influences its friends"],
       ["Link (≈250)", "A friendship",
        "Channel through which information and norms are transmitted. The network is fixed during a run"],
       ["Product: orange (0) and blue (1)", "Two competing brands",
        "Defined by global parameters (`quality_orange`, `quality_blue`) that can be changed during a run"]],
      [30, 34, 58])
P("Time is discrete. In each tick all consumers first execute the consumption decision and "
  "then update their utilities. Within each phase the consumers are activated in random order "
  "(Listing 1). The outputs are recorded at the end of each tick.")
listing("consumer_choice/model.py", [("def step(self)", "datacollector.collect(self)")], "Schedule of one tick")

H2("2.2   State Variables of the Consumer")
P("Each consumer holds five state variables (Table 2, Listing 2), and a position used only to "
  "build the network (Sect. 2.3). The variable `satisfaction` links "
  "product quality to habit: it is the probability that the consumer keeps its current "
  "product without considering the alternative. The variable `personal_experience` is the "
  "memory of the consumer; it is updated only for the product in use and may therefore become "
  "outdated when the quality of the other product changes.")
table("State variables of the consumer.",
      ["Variable", "Values", "Role", "Consequence"],
      [["`choice`", "0 (orange) or 1 (blue)", "Product in use",
        "Counts for the market share; it is what friends observe when they compute the norm"],
       ["`satisfaction`", "log10(1 + 9q) of the product in use", "Probability of keeping the product (habit)",
        "With q = 0.5 it equals 0.74. Higher quality produces stronger habits and less attention to the alternative"],
       ["`personal_experience`", "Remembered quality of each product, initially [0, 0]",
        "Memory of the consumer", "A product never used is valued at 0; a change in the quality of the unused product is not perceived"],
       ["`utilities`", "Total utility of each product", "Basis of the choice when the consumer neither keeps nor explores",
        "Combines norm and product utility (Sect. 2.4)"],
       ["`friends`", "List of linked consumers", "Social network of the consumer",
        "Source of information and of the norm"]],
      [32, 27, 27, 36])
listing("consumer_choice/model.py", [("class Consumer(", "self.friends = []")], "Consumer state variables")

H2("2.3   Social Network")
P("Each consumer is placed at a random position in a 51 × 51 space. Links are created one at "
  "a time: a random consumer is linked to the nearest consumer to which it is not yet linked, "
  "until the number of links equals `avg_connections` × `num_consumers` / 2 (250 in the "
  "baseline). The position is used only for this purpose (Listing 3). Because links connect "
  "nearby consumers, the network is local: friends of a consumer tend to be friends of each "
  "other. This property favours the formation of clusters of consumers that use the same "
  "product.")
listing("consumer_choice/model.py", [("def build_network", "links += 1")], "Construction of the social network")

H2("2.4   Actions of the Consumer")
H3("Consumption Decision.", "In each tick the consumer applies three rules in sequence "
   "(Listing 4). First, with probability equal to `satisfaction`, it keeps the current product "
   "(habit). Second, if it did not keep the product, it switches to the other product with "
   "probability `exploration`. Third, if it neither kept nor explored, it chooses the product "
   "with the higher utility; if both utilities are equal, it keeps the current product. "
   "Satisfaction is then recomputed for the product in use. In the baseline the probability "
   "of exploring in a given tick is (1 − 0.74) × 0.01 ≈ 0.003.")
listing("consumer_choice/model.py", [("def consume(self)", "self.satisfaction = self.model.satisfaction(self.choice)"), ("def satisfaction(self, choice)", "return math.log10")], "Consumption decision and satisfaction")
H3("Utility Update.", "After the consumption phase, each consumer computes the utility of "
   "each product (Listing 5) in four steps: (1) the experience of the product in use is set to "
   "its current quality; (2) with probability `information_exchange`, and only if at least one "
   "friend uses the product, the consumer receives the current quality of that product; "
   "(3) the norm-based utility is the share of friends who use the product; (4) the total "
   "utility is")
p = doc.add_paragraph(style="Normal"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_runs(p, "*U* = *w* · *N* + (1 − *w*) · (*E* + *I*) / 2 ,")
P("where *w* is `norm_influence`, *N* the share of friends using the product, *E* the "
  "remembered quality and *I* the quality reported by friends (0 if no information was "
  "received). Without information from friends, the product utility is half of the "
  "remembered quality. As a result, information exchange favours the product that friends "
  "already use, in the same direction as the norm.")
listing("consumer_choice/model.py", [("def update_utilities", "(1 - m.norm_influence) * prod_util)")], "Utility update")
P("Table 3 summarises the actions and their consequences.")
table("Actions of the consumer, in the order in which they are executed.",
      ["Action", "When", "Consequence"],
      [["Keep the product (habit)", "With probability `satisfaction`",
        "No change. The consumer does not consider the alternative, so improvements of the other product are not perceived"],
       ["Explore", "If not kept, with probability `exploration`",
        "Switches to the other product and learns its actual quality in the next update. It is the only way to discover an improvement of a product that no friend uses"],
       ["Choose by utility", "If neither kept nor explored, and utilities differ",
        "Selects the product with the higher utility; in a tie it keeps the current product"],
       ["Update experience", "Every tick, for the product in use", "Memory of that product equals its current quality"],
       ["Ask friends", "For each product, with probability `information_exchange`, if a friend uses it",
        "Receives the current quality. Reinforces the product used by friends"],
       ["Observe the norm", "Every tick", "Pulls utility towards the local majority; produces clusters and can produce lock-in"]],
      [28, 38, 56])
P("Taken together, habit slows down change, exploration is the only source of new "
  "information about an unused product, and information exchange and the norm reinforce the "
  "product that is already used by friends. These mechanisms produce the three patterns "
  "described by the authors of the model: clustering of consumption, amplification of small "
  "differences in quality and social lock-in [1].")

H2("2.5   Global Parameters and Baseline")
P("The model has eight parameters (Listing 6). The baseline uses the default values of the "
  "original model file (Table 4).")
listing("consumer_choice/model.py", [("def __init__(self, num_consumers", "self.quality = [quality_orange")], "Model parameters and their baseline values")
table("Global parameters, baseline values and role. The last column summarises the effect "
      "measured in Sect. 3.3; parameters marked “not tested” were kept at the baseline value.",
      ["Parameter", "Baseline", "Role", "Effect of increasing the value"],
      [["`num_consumers`", "100", "Size of the market", "Not tested"],
       ["`avg_connections`", "5", "Average number of friends", "Not tested"],
       ["`initial_pct_orange`", "50", "Initial share of orange consumers", "Not tested"],
       ["`norm_influence`", "0.15", "Weight of the norm against product utility",
        f"From 0 to 0.1, mean clustering rises from {f2(SN[0.0]['clustering_mean'])} to {f2(SN[0.1]['clustering_mean'])}; above 0.1 it changes little"],
       ["`information_exchange`", "0.05", "Probability of learning from friends", "Not tested"],
       ["`exploration`", "0.01", "Probability of trying the other product",
        f"At 0.2, mean clustering falls to {f2(SE[0.2]['clustering_mean'])} and the dispersion of market shares decreases"],
       ["`quality_orange`", "0.5", "Quality of the orange product", "Not tested (reference value)"],
       ["`quality_blue`", "0.5", "Quality of the blue product",
        f"A difference of 0.05 gives the better product {round((1 - SQ[0.55]['orange_mean']) * 100)}–{round(SQ[0.45]['orange_mean'] * 100)}% of the market (mean)"]],
      [33, 15, 31, 43])

H2("2.6   Model Outputs")
P("The original model reports two plots: the market shares of the two products and the "
  "clustering of consumers. The implementation records the same three quantities at every "
  "tick with a Mesa `DataCollector` (Listing 7, Table 5). The analysis in Section 3 uses only "
  "these outputs.")
listing("consumer_choice/model.py", [("self.datacollector = mesa.DataCollector", "self.datacollector.collect(self)"), ("def count_choice", "return total / self.num_consumers")], "Model outputs")
table("Outputs of the model.",
      ["Output", "Definition", "Interpretation"],
      [["Orange consumers", "Number of consumers with `choice` = 0", "Market share of the orange product (out of 100)"],
       ["Blue consumers", "Number of consumers with `choice` = 1", "Market share of the blue product (out of 100)"],
       ["Clustering", "Mean, over consumers, of the share of friends that use the same product",
        "1: all friends use the same product; ≈0.5: choices independent of friends"]],
      [27, 50, 45])

H2("2.7   Notes on the Implementation")
P("Three differences between the description in the Info tab of the original model and its "
  "code were found; the implementation follows the code:")
bullets(["The Info tab states a default `exploration` of 0.05; the model file uses 0.01.",
         "The Info tab states that friends may share outdated information. In the code, a friend "
         "always reports the current quality of the product.",
         "The Info tab states that less satisfied consumers seek more information. In the code, "
         "the probability of information exchange does not depend on satisfaction; satisfaction "
         "only determines whether the utilities are used in the decision."])

# ================================================================ 3 Results
H1("3   Baseline Simulation and Statistical Analysis")
H2("3.1   Experimental Setup")
P(f"The baseline was replicated {RUNS} times with different random seeds (0 to {RUNS - 1}), "
  "for 2000 ticks each, with the outputs recorded at every tick. The statistics are reported "
  "at tick 1000, and ticks 1500 and 2000 are used to check the horizon. For the sensitivity "
  "analysis, three parameters were varied one at a time, with 100 replications per value and "
  "1000 ticks per run (Table 6). The same 100 seeds were used for every value of a parameter "
  "(common random numbers), so that differences between values are not due to different "
  "random draws. The experiments use the Mesa `batch_run` function (Listing 8).")
table("Experimental design.",
      ["Experiment", "Parameter values", "Runs per value", "Ticks"],
      [["Baseline", "Table 4", str(RUNS), "2000"],
       ["Norm", "`norm_influence` = 0, 0.1, …, 1.0", "100", "1000"],
       ["Quality", "`quality_blue` = 0.3, 0.4, 0.45, 0.5, 0.55, 0.6, 0.7", "100", "1000"],
       ["Exploration", "`exploration` = 0, 0.01, 0.05, 0.1, 0.2", "100", "1000"]],
      [22, 62, 20, 18])
listing("experiments/run_experiments.py", [("BASELINE = dict(", "quality_orange=0.5, quality_blue=0.5)"), ("results = batch_run(", "display_progress=False)")], "Baseline values and batch run")
P("For each output we report the mean, the standard deviation (SD), the 95% confidence "
  "interval of the mean (Student’s *t*), the median and the 5th and 95th percentiles "
  "(P5, P95). Market shares are expressed as fractions of the 100 consumers.")
H3("Number of Runs.", "Figure 2 shows the mean of each output at tick 1000 as a function of the "
   f"number of runs. With 100 runs, the half-width of the 95% confidence interval is "
   f"{f3(CI100['orange_share'])} for the orange share and {f3(CI100['clustering'])} for "
   f"clustering. With {RUNS} runs it is {f3(CIALL['orange_share'])} and "
   f"{f3(CIALL['clustering'])}. The market share requires more runs than clustering because "
   "its variance between runs is "
   f"about {round((Bo['sd'] / Bc['sd']) ** 2)} times larger. We therefore use "
   f"{RUNS} runs for the baseline and 100 runs per value in the sensitivity analysis, where "
   "the differences between parameter values are larger than this precision.")
H3("Horizon.", "Table 8 shows the outputs at several ticks. Mean clustering increases by "
   f"{f3(HZ[1000]['clustering_mean'] - HZ[500]['clustering_mean'])} between ticks 500 and 1000 and by "
   f"{f3(HZ[2000]['clustering_mean'] - HZ[1000]['clustering_mean'])} between ticks 1000 and 2000. "
   "The SD of the orange share keeps increasing "
   f"({f2(HZ[500]['orange_share_sd'])}, {f2(HZ[1000]['orange_share_sd'])}, {f2(HZ[2000]['orange_share_sd'])} "
   "at ticks 500, 1000 and 2000). The model does not reach a stationary distribution of market "
   "shares within 2000 ticks: individual markets continue to move away from the initial "
   "50/50 split. Tick 1000 is used as the reference horizon because the change of mean "
   "clustering after it is small and the shares have already separated. Statistics at a later "
   "tick would show a larger dispersion of shares; this must be considered when the extended "
   "model is compared with the baseline.")

H2("3.2   Baseline Results")
P(f"Table 7 reports the statistics of the three outputs at tick 1000 over the {RUNS} runs, and "
  "Fig. 1 shows their evolution.")
table(f"Baseline outputs at tick 1000 ({RUNS} runs).",
      ["Output", "Mean", "SD", "95% CI", "Median", "P5", "P95"],
      [[name, f3(x["mean"]), f3(x["sd"]), f"{f3(x['ci95_low'])}–{f3(x['ci95_high'])}",
        f3(x["median"]), f2(x["p5"]), f2(x["p95"])]
       for name, x in [("Orange market share", Bo), ("Blue market share", Bb), ("Clustering", Bc)]],
      [30, 12, 12, 24, 14, 15, 15], align="lrrrrrr")
figure(f"{FIG}/fig1_baseline.png",
       "Baseline. (a) Orange market share in three runs (the 5th, 50th and 95th percentile of the "
       f"final share). (b) Mean clustering and mean orange share over {RUNS} runs, with the "
       "band between P5 and P95. (c) Distribution of the orange share at tick 1000.")
P("Three results characterise the baseline:")
bullets([f"*Each market becomes segmented.* Mean clustering rises from {f2(HZ[0]['clustering_mean'])} "
         f"at tick 0 to {f2(Bc['mean'])} at tick 1000 (95% CI {f3(Bc['ci95_low'])}–{f3(Bc['ci95_high'])}). "
         f"On average, {round(Bc['mean'] * 100)}% of the friends of a consumer use the same "
         "product. The variation between runs is small (SD "
         f"{f3(Bc['sd'])}), so this result does not depend on the random seed.",
         f"*The market outcome is not predictable.* The mean orange share is {f2(Bo['mean'])}, as "
         "expected from the symmetry of the baseline (equal qualities and initial shares), but "
         f"the SD between runs is {f2(Bo['sd'])}. In 90% of the runs the orange share lies between "
         f"{f2(Bo['p5'])} and {f2(Bo['p95'])}. Orange holds the majority in {OUTC['pct_runs_orange_majority']:.1f}% "
         f"of the runs and blue in {OUTC['pct_runs_blue_majority']:.1f}% (ties: {OUTC['pct_runs_tie']:.1f}%). Only "
         f"{OUTC['pct_runs_share_between_40_60']:.0f}% of the runs end with shares between 0.40 and 0.60, and "
         f"{OUTC['pct_runs_one_product_ge_90pct']:.1f}% end with one product at 90% or more.",
         "*The mean share is not informative; the dispersion is.* With identical products, "
         "the result of a run depends on early random events that are amplified by the norm and "
         "by information exchange. The relevant statistic for the baseline is therefore the "
         "distribution of shares (Fig. 1c), not its mean."])
table("Mean and SD of the outputs over time (baseline).",
      ["Tick", "Clustering (mean)", "Clustering (SD)", "Orange share (mean)", "Orange share (SD)"],
      [[str(t), f3(HZ[t]["clustering_mean"]), f3(HZ[t]["clustering_sd"]),
        f3(HZ[t]["orange_share_mean"]), f3(HZ[t]["orange_share_sd"])] for t in [0, 100, 250, 500, 750, 1000, 1500, 2000]],
      [14, 27, 27, 27, 27], align="lrrrr")
figure(f"{FIG}/fig2_number_of_runs.png",
       "Mean of the outputs at tick 1000 as a function of the number of runs, with the 95% "
       "confidence interval. (a) Orange market share. (b) Clustering.")

H2("3.3   Sensitivity Analysis")
P("The sensitivity analysis varies one parameter at a time, with all other parameters at "
  "the baseline. Figures 3 to 5 show the distribution of the orange share and the mean "
  "clustering at tick 1000, and Tables 9 to 11 report the statistics.")
H3("Social Norm.", f"Without norm (`norm_influence` = 0), mean clustering is "
   f"{f2(SN[0.0]['clustering_mean'])}, close to the value of independent choices, and the orange "
   f"share stays near 0.5 (SD {f2(SN[0.0]['orange_sd'])}). With `norm_influence` = 0.1, mean "
   f"clustering is {f2(SN[0.1]['clustering_mean'])} and the SD of the orange share is "
   f"{f2(SN[0.1]['orange_sd'])}. Above 0.1 both values change little (clustering between "
   f"{f2(min(SN[v]['clustering_mean'] for v in SN if v > 0))} and "
   f"{f2(max(SN[v]['clustering_mean'] for v in SN if v > 0))}). The effect of the norm is a "
   "threshold, not a gradual effect. The reason is visible in the utility function. With "
   "both qualities equal to *q*, the product utility of a consumer can only take the values 0, "
   "*q*/2 and *q*, so the difference between the two products is either 0 or at least *q*/2. "
   "The norm can reverse a non-zero difference only if *w* / (1 − *w*) > *q*/2, that is, "
   "*w* > *q* / (2 + *q*) = 0.2 for *q* = 0.5. For *w* ≤ 0.2 the norm acts only when the product "
   "utilities are equal. This explains why the results "
   "for *w* = 0.1 and *w* = 0.2 are identical under common random numbers (Table 9), and why "
   "a small positive norm is sufficient to coordinate the choices of friends.")
figure(f"{FIG}/fig3_sensitivity_norm_influence.png",
       "Effect of `norm_influence` (100 runs per value). (a) Distribution of the orange share "
       "at tick 1000. (b) Mean clustering at tick 1000 with 95% confidence interval.")
table("Outputs at tick 1000 as a function of `norm_influence` (100 runs per value).",
      ["`norm_influence`", "Orange share (mean)", "Orange share (SD)", "Clustering (mean)", "Clustering 95% CI"],
      [[f"{v:g}", f3(r["orange_mean"]), f3(r["orange_sd"]), f3(r["clustering_mean"]),
        f"{f3(r['clustering_ci95_low'])}–{f3(r['clustering_ci95_high'])}"] for v, r in sorted(SN.items())],
      [26, 24, 22, 23, 27], align="lrrrr")
H3("Product Quality.", "With `quality_orange` fixed at 0.5, a small difference in quality "
   f"determines the market. With `quality_blue` = 0.45, the mean orange share is "
   f"{f2(SQ[0.45]['orange_mean'])}; with 0.55, it is {f2(SQ[0.55]['orange_mean'])}, which gives blue a "
   f"mean share of {f2(1 - SQ[0.55]['orange_mean'])}. With a difference of 0.1 or more, the better product "
   f"holds at least {min(SQ[0.4]['orange_mean'], 1 - SQ[0.6]['orange_mean']) * 100:.1f}% of the market "
   "in the mean. Mean clustering is higher when qualities differ "
   f"({f2(SQ[0.3]['clustering_mean'])} at 0.3 and {f2(SQ[0.7]['clustering_mean'])} at 0.7) than when they are "
   f"equal ({f2(SQ[0.5]['clustering_mean'])}), because most consumers then use the same product. "
   "The dispersion of the outcome is largest when the products are equal (SD "
   f"{f2(SQ[0.5]['orange_sd'])}) and falls to {f3(max(SQ[v]['orange_sd'] for v in (0.3, 0.4, 0.6, 0.7)))} or less when the difference "
   "is 0.1 or more.")
figure(f"{FIG}/fig4_sensitivity_quality_blue.png",
       "Effect of `quality_blue`, with `quality_orange` = 0.5 (100 runs per value). (a) "
       "Distribution of the orange share at tick 1000. (b) Mean clustering at tick 1000 with "
       "95% confidence interval.")
table("Outputs at tick 1000 as a function of `quality_blue` (`quality_orange` = 0.5; 100 runs per value).",
      ["`quality_blue`", "Orange share (mean)", "Orange share (SD)", "Clustering (mean)", "Clustering 95% CI"],
      [[f"{v:g}", f3(r["orange_mean"]), f3(r["orange_sd"]), f3(r["clustering_mean"]),
        f"{f3(r['clustering_ci95_low'])}–{f3(r['clustering_ci95_high'])}"] for v, r in sorted(SQ.items())],
      [26, 24, 22, 23, 27], align="lrrrr")
H3("Exploration.", "Mean clustering decreases as exploration increases, from "
   f"{f2(SE[0.01]['clustering_mean'])} at 0.01 to {f2(SE[0.05]['clustering_mean'])} at 0.05 and "
   f"{f2(SE[0.2]['clustering_mean'])} at 0.2, because random switches break the agreement between "
   "friends. The dispersion of the orange share is not monotonic: its SD is "
   f"{f2(SE[0.0]['orange_sd'])} without exploration, {f2(SE[0.05]['orange_sd'])} at 0.05 and "
   f"{f2(SE[0.2]['orange_sd'])} at 0.2. A possible interpretation is that, without exploration, few consumers change product and "
   "the shares move slowly; moderate exploration increases the number of changes and lets "
   "clusters grow; high exploration randomises choices and keeps the shares close to 0.5.")
figure(f"{FIG}/fig5_sensitivity_exploration.png",
       "Effect of `exploration` (100 runs per value). (a) Distribution of the orange share at "
       "tick 1000. (b) Mean clustering at tick 1000 with 95% confidence interval.")
table("Outputs at tick 1000 as a function of `exploration` (100 runs per value).",
      ["`exploration`", "Orange share (mean)", "Orange share (SD)", "Clustering (mean)", "Clustering 95% CI"],
      [[f"{v:g}", f3(r["orange_mean"]), f3(r["orange_sd"]), f3(r["clustering_mean"]),
        f"{f3(r['clustering_ci95_low'])}–{f3(r['clustering_ci95_high'])}"] for v, r in sorted(SE.items())],
      [26, 24, 22, 23, 27], align="lrrrr")

H2("3.4   Summary of Findings")
P("The baseline and the sensitivity analysis lead to four conclusions that will serve as "
  "reference for the extended model:")
bullets([f"Social influence segments the market. All tested positive values of `norm_influence` "
         f"(0.1 to 1.0) lead to a mean clustering of about "
         f"{f2(Bc['mean'])}; without norm, clustering stays at {f2(SN[0.0]['clustering_mean'])}.",
         "With equal products, the outcome of a market is path dependent: the same rules produce "
         f"markets dominated by either product (SD of the share {f2(Bo['sd'])}). The baseline must be "
         "described by the distribution of outcomes, not by the mean.",
         "Small differences in quality are amplified: a difference of 0.05 gives the better "
         f"product about {round((1 - SQ[0.55]['orange_mean']) * 100)}–{round(SQ[0.45]['orange_mean'] * 100)}% of the market.",
         "Exploration is the only mechanism that brings new information into a cluster; it reduces "
         "clustering but also determines how fast markets can change."])

# ================================================================ 4 Next steps
H1("4   Next Step: Fitting the Model to a Real Market")
P("In the next phase, the model will be fitted to the entry of a new operator in the "
  "Portuguese mobile telecommunications market. DIGI entered the market at the end of 2024 "
  "with prices below those of MEO, NOS and Vodafone. According to data from the national "
  "regulator ANACOM, it held 2.8% of bundled-service customers in the first half of 2025 [8].")
P("*Problem.* Which combination of price, network quality and trial campaigns allows a new "
  "operator to cross the tipping point, and how long does it take? The same question can be "
  "asked from the point of view of the incumbents: which response limits the loss of market "
  "share?")
table("Correspondence between the original model and the proposed model.",
      ["Original model", "Proposed model"],
      [["2 products", "4 operators (MEO, NOS, Vodafone, DIGI); DIGI enters at tick T"],
       ["Quality", "Network and service quality (ANACOM indicators or network tests)"],
       ["—", "Price, as a new attribute of utility"],
       ["Habit (satisfaction)", "Habit and contract lock-in (24-month contracts: a consumer can switch only when the contract ends)"],
       ["Social norm", "Family plans and recommendations of friends"],
       ["Exploration", "Trial periods, number portability, promotions"]],
      [40, 82])
P("*Steps.*")
bullets(["Generalise the code from two to *N* products (lists instead of `1 − choice`).",
         "Add price and a heterogeneous sensitivity to price, for example "
         "*U* = *a* · quality − *b* · price + norm, with *b* specific to each consumer.",
         "Add the contract (a counter per agent) and the entry of new operators during a run.",
         "Calibrate the weights with the ANACOM market shares before the entry of DIGI.",
         "Validate by comparing the simulated adoption curve of DIGI with the quarterly data.",
         "Experiments: grid of DIGI price × quality investment × intensity of trial campaigns; "
         "measure the DIGI share after two years and the time to reach 10%.",
         "Learning agents (task 5): replace the fixed exploration rate by consumers that learn as "
         "in a multi-armed bandit (ε-greedy or UCB) [9], or let operators learn prices with Q-learning.",
         "Compare the new model with the baseline of this report (task 6), in particular the "
         "clustering and the dependence on early random events observed for "
         "`norm_influence` > 0."])
P("This case fits the mechanisms of the model: switching between operators is infrequent "
  "(habit and contracts), family plans and recommendations act as a norm, and trial offers "
  "and portability correspond to exploration. Public market-share data allow calibration and "
  "validation.")

# ================================================================ References
H1("References")
refs = [
    "Jager, W., Wang, S.: NetLogo Consumer Choice model. Center for Connected Learning and "
    "Computer-Based Modeling, Northwestern University, Evanston, IL (2024). "
    "http://ccl.northwestern.edu/netlogo/models/ConsumerChoice",
    "Wilensky, U.: NetLogo. Center for Connected Learning and Computer-Based Modeling, "
    "Northwestern University, Evanston, IL (1999). http://ccl.northwestern.edu/netlogo/",
    "Janssen, M.A., Jager, W.: Fashions, Habits and Changing Preferences: Simulation of "
    "Psychological Factors Affecting Market Dynamics. Journal of Economic Psychology 22(6), "
    "745–772 (2001)",
    "Cialdini, R.B., Goldstein, N.J.: Social Influence: Compliance and Conformity. Annual "
    "Review of Psychology 55, 591–621 (2004)",
    "Nyborg, K., et al.: Social Norms as Solutions. Science 354(6308), 42–43 (2016)",
    "ter Hoeven, E., Kwakkel, J., Hess, V., Pike, T., Wang, B., Rht, Kazil, J.: Mesa 3: "
    "Agent-Based Modeling with Python in 2025. Journal of Open Source Software 10(107), 7668 "
    "(2025). https://doi.org/10.21105/joss.07668",
    "Grimm, V., et al.: The ODD Protocol for Describing Agent-Based and Other Simulation "
    "Models: A Second Update to Improve Clarity, Replication, and Structural Realism. Journal "
    "of Artificial Societies and Social Simulation 23(2), 7 (2020)",
    "4gnews: DIGI começa 2026 a mostrar à MEO, NOS e Vodafone como se faz (ANACOM data, first "
    "half of 2025). https://4gnews.pt/digi-comeca-2026-a-mostrar-a-meo-nos-e-vodafone-como-se-faz/",
    "Sutton, R.S., Barto, A.G.: Reinforcement Learning: An Introduction, 2nd edn. MIT Press, "
    "Cambridge, MA (2018)",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph(style="reference")
    p.add_run(f"{i}.\t{ref}")

doc.save(OUT)

# remove the template's sample comments and footnotes (no longer referenced)
import shutil, zipfile, tempfile  # noqa: E402
tmp = tempfile.mktemp(suffix=".docx")
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/comments.xml":
            data = re.sub(rb"<w:comment\b.*?</w:comment>", b"", data, flags=re.S)
        elif item.filename == "word/numbering.xml":
            # schema order inside <w:lvl>: numFmt before pStyle (template converted from .doc)
            data = re.sub(rb"(<w:pStyle [^>]*/>)(<w:numFmt [^>]*/>)", rb"\2\1", data)
        elif item.filename == "word/footnotes.xml":
            data = re.sub(rb'<w:footnote (?:(?!w:type=)[^>])*>.*?</w:footnote>', b"", data, flags=re.S)
        zout.writestr(item, data)
shutil.move(tmp, OUT)
print("saved", OUT, counters)
